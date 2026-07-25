"""Financial spreading: parse CSV/XLSX, normalize, variance vs stated figures."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from packages.schemas.case import VarianceFinding

# Canonical metric keys
CANONICAL = {
    "gross_revenue": ["gross_revenue", "annual_revenue", "revenue", "gross_sales", "sales"],
    "noi": ["noi", "net_operating_income"],
    "avg_monthly_deposits": [
        "avg_monthly_deposits",
        "average_monthly_deposits",
        "monthly_deposits",
        "avg_deposits",
    ],
    "stated_income": [
        "stated_income",
        "monthly_income",
        "stated_monthly_income",
        "annual_salary",
        "salary",
    ],
    "ebitda": ["ebitda"],
}

THRESHOLDS = {
    "sba_7a": {"gross_revenue": 0.10, "stated_income": 0.15, "default": 0.15},
    "cre_acquisition": {"noi": 0.10, "gross_revenue": 0.15, "default": 0.15},
    "specialty_mortgage_bank_statement": {
        "avg_monthly_deposits": 0.15,
        "stated_income": 0.15,
        "default": 0.15,
    },
}


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", h.strip().lower()).strip("_")


def _to_float(val: Any) -> float | None:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace("$", "").replace("%", "")
    try:
        return float(s)
    except ValueError:
        return None


def map_to_canonical(header: str) -> str | None:
    key = _norm_header(header)
    for canon, aliases in CANONICAL.items():
        if key == canon or key in aliases:
            return canon
    return None


def parse_tabular(filename: str, data: bytes) -> dict[str, float]:
    """Return canonical_metric -> value from CSV or XLSX."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xltx")):
        return _parse_xlsx(data)
    return _parse_csv(data)


def _parse_csv(data: bytes) -> dict[str, float]:
    text = data.decode("utf-8-sig", errors="ignore")
    # Two-column metric,value format
    reader2 = csv.reader(io.StringIO(text))
    rows2 = [r for r in reader2 if r]
    if rows2 and len(rows2[0]) >= 2:
        h0, h1 = _norm_header(rows2[0][0]), _norm_header(rows2[0][1])
        if h0 in ("metric", "key", "name", "field") or h1 in ("value", "amount"):
            out: dict[str, float] = {}
            for row in rows2[1:]:
                if len(row) < 2:
                    continue
                canon = map_to_canonical(row[0])
                num = _to_float(row[1])
                if canon and num is not None:
                    out[canon] = num
            if out:
                return out

    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        out = {}
        for row in rows2:
            if len(row) < 2:
                continue
            canon = map_to_canonical(row[0])
            num = _to_float(row[1])
            if canon and num is not None:
                out[canon] = num
        return out
    out = {}
    for row in rows:
        for h, v in row.items():
            if h is None:
                continue
            canon = map_to_canonical(h)
            num = _to_float(v)
            if canon and num is not None:
                out[canon] = num
    return out


def _parse_xlsx(data: bytes) -> dict[str, float]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    headers = [str(h or "") for h in rows[0]]
    out: dict[str, float] = {}
    for row in rows[1:]:
        for h, v in zip(headers, row, strict=False):
            canon = map_to_canonical(h)
            num = _to_float(v)
            if canon and num is not None:
                out[canon] = num
    # Also support metric/value pairs in col A/B
    if not out:
        for row in rows:
            if not row or len(row) < 2:
                continue
            canon = map_to_canonical(str(row[0] or ""))
            num = _to_float(row[1])
            if canon and num is not None:
                out[canon] = num
    return out


def stated_from_metadata(metadata: dict[str, Any]) -> dict[str, float]:
    out: dict[str, float] = {}
    mapping = {
        "gross_revenue": ["gross_revenue", "annual_revenue", "stated_revenue"],
        "stated_income": [
            "stated_income",
            "monthly_income",
            "stated_monthly_income",
            "annual_salary",
        ],
        "avg_monthly_deposits": ["avg_monthly_deposits", "average_monthly_deposits"],
        "noi": ["noi", "net_operating_income"],
    }
    for canon, keys in mapping.items():
        for k in keys:
            if k in metadata and metadata[k] is not None:
                num = _to_float(metadata[k])
                if num is not None:
                    out[canon] = num
                    break
    return out


def compute_variances(
    spread: dict[str, float],
    stated: dict[str, float],
    vertical: str,
) -> list[VarianceFinding]:
    thresholds = THRESHOLDS.get(vertical, THRESHOLDS["sba_7a"])
    findings: list[VarianceFinding] = []
    for metric, spread_val in spread.items():
        stated_val = stated.get(metric)
        if stated_val is None or stated_val == 0:
            continue
        thr = float(thresholds.get(metric, thresholds.get("default", 0.15)))
        variance_pct = abs(spread_val - stated_val) / abs(stated_val)
        if variance_pct > thr:
            severity = "high" if variance_pct > thr * 2 else "medium"
            findings.append(
                VarianceFinding(
                    metric=metric,
                    stated=stated_val,
                    spread=spread_val,
                    variance_pct=round(variance_pct, 4),
                    threshold_pct=thr,
                    severity=severity,
                    summary=(
                        f"{metric}: stated {stated_val:,.2f} vs spread {spread_val:,.2f} "
                        f"({variance_pct:.1%} > {thr:.0%} threshold)"
                    ),
                )
            )
    return findings
